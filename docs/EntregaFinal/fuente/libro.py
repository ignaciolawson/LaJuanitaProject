"""Generador del libro de la entrega final.

Arma el .docx con el formato pedido por la cátedra (Calibri 11, interlineado 1,5,
texto justificado, títulos y tablas en APA 7) y preparado para imprimir a doble faz
y encuadernar en tapa dura: A4, márgenes espejo con más margen del lado del lomo,
cada parte arranca en página impar.

El contenido vive en partes/*.py, un módulo por parte del libro. Este archivo solo
sabe de formato.

    python docs/EntregaFinal/fuente/libro.py
    pwsh docs/EntregaFinal/fuente/exportar.ps1     # índice + PDF, con Word
"""

import importlib
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[2]
SALIDA = AQUI.parent / "La-Juanita-Studio-Entrega-Final.docx"

FUENTE = "Calibri"
GRIS = RGBColor(0x59, 0x59, 0x59)

# A4. El margen interior (lomo) es más grande: la tapa dura cosida se come ~1 cm.
ANCHO_PAGINA = Cm(21.0)
ALTO_PAGINA = Cm(29.7)
MARGEN_INTERIOR = Cm(3.5)
MARGEN_EXTERIOR = Cm(2.5)
MARGEN_VERTICAL = Cm(2.54)
ANCHO_TEXTO = ANCHO_PAGINA - MARGEN_INTERIOR - MARGEN_EXTERIOR  # 15 cm

TITULO_CORTO = "La Juanita Studio · Sistema de Gestión"

# Las partes, en orden. Cada una es un módulo de partes/ con escribir(libro).
PARTES = [
    ("I", "Estudio de Prefactibilidad", "p1_prefactibilidad"),
    ("II", "Informe de Relevamiento", "p2_relevamiento"),
    ("III", "Propuesta Técnica", "p3_propuesta_tecnica"),
    ("IV", "Propuesta Comercial", "p4_propuesta_comercial"),
    ("V", "Diagramas de Base de Datos", "p5_diagramas_bd"),
    ("VI", "Planilla de Tiempos", "p6_planilla_tiempos"),
    ("VII", "Diagrama de Gantt", "p7_gantt"),
    ("VIII", "Diccionario de Datos", "p8_diccionario"),
]


# --------------------------------------------------------------------------- XML


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def _fuente(rpr_owner, nombre=FUENTE):
    rpr = rpr_owner.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = _el("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(k), nombre)
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(k)) is not None:
            del fonts.attrib[qn(k)]


def _campo(parrafo, instruccion, texto=""):
    """Un campo de Word (PAGE, TOC…). Word lo calcula al abrir o al exportar."""
    r = parrafo.add_run()
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "begin", "w:dirty": "true"}))
    r = parrafo.add_run()
    it = _el("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instruccion} "
    r._r.append(it)
    r = parrafo.add_run()
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    parrafo.add_run(texto)
    r = parrafo.add_run()
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))


def _numeracion(seccion, formato, inicio=None):
    sectpr = seccion._sectPr
    for viejo in sectpr.findall(qn("w:pgNumType")):
        sectpr.remove(viejo)
    pg = _el("w:pgNumType", **{"w:fmt": formato})
    if inicio is not None:
        pg.set(qn("w:start"), str(inicio))
    sectpr.append(pg)


def _borde_inferior(parrafo, color="808080"):
    ppr = parrafo._p.get_or_add_pPr()
    bdr = _el("w:pBdr")
    bdr.append(_el("w:bottom", **{"w:val": "single", "w:sz": 4, "w:space": 4, "w:color": color}))
    ppr.append(bdr)


# ------------------------------------------------------------------- texto rico

_MARCAS = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def escribir_rico(parrafo, texto, tamano=None):
    """**negrita** y *cursiva*, nada más."""
    for trozo in _MARCAS.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**"):
            r = parrafo.add_run(trozo[2:-2])
            r.bold = True
        elif trozo.startswith("*"):
            r = parrafo.add_run(trozo[1:-1])
            r.italic = True
        else:
            r = parrafo.add_run(trozo)
        if tamano:
            r.font.size = Pt(tamano)
    return parrafo


# ---------------------------------------------------------------------- el libro


class Libro:
    def __init__(self):
        self.doc = Document()
        self.figuras = 0
        self.tablas = 0
        self._estilos()
        s = self.doc.sections[0]
        self._geometria(s)
        settings = self.doc.settings.element
        settings.append(_el("w:mirrorMargins"))
        settings.append(_el("w:evenAndOddHeaders"))
        # Que Word recalcule el índice al abrir, por si se abre sin pasar por exportar.ps1.
        settings.append(_el("w:updateFields", **{"w:val": "true"}))

    # -- estilos -------------------------------------------------------------

    def _estilos(self):
        st = self.doc.styles

        normal = st["Normal"]
        normal.font.name = FUENTE
        normal.font.size = Pt(11)
        _fuente(normal.element)
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.widow_control = True

        # Cuerpo: sangría de primera línea de 1,27 cm, como pide APA.
        cuerpo = st.add_style("Cuerpo", WD_STYLE_TYPE.PARAGRAPH)
        cuerpo.base_style = normal
        cuerpo.paragraph_format.first_line_indent = Cm(1.27)

        # Títulos APA 7: nivel 1 centrado y negrita; nivel 2 a la izquierda y negrita;
        # nivel 3 a la izquierda, negrita y cursiva. Mismo tamaño que el texto.
        # Heading 1 queda para las partes del libro (I, II, III…), que no son APA.
        def titulo(nombre, tam, alineacion, cursiva=False, antes=12, despues=6):
            h = st[nombre]
            h.font.name = FUENTE
            _fuente(h.element)
            h.font.size = Pt(tam)
            h.font.bold = True
            h.font.italic = cursiva
            h.font.color.rgb = RGBColor(0, 0, 0)
            f = h.paragraph_format
            f.alignment = alineacion
            f.space_before = Pt(antes)
            f.space_after = Pt(despues)
            f.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            f.keep_with_next = True
            f.first_line_indent = Cm(0)
            return h

        titulo("Heading 1", 26, WD_ALIGN_PARAGRAPH.CENTER, antes=0, despues=0)
        titulo("Heading 2", 11, WD_ALIGN_PARAGRAPH.CENTER, antes=18, despues=6)
        titulo("Heading 3", 11, WD_ALIGN_PARAGRAPH.LEFT, antes=12, despues=3)
        titulo("Heading 4", 11, WD_ALIGN_PARAGRAPH.LEFT, cursiva=True, antes=12, despues=3)

        for nivel in (1, 2, 3):
            t = st[f"TOC {nivel}"] if f"TOC {nivel}" in [s.name for s in st] else st.add_style(
                f"TOC {nivel}", WD_STYLE_TYPE.PARAGRAPH)
            t.base_style = normal
            t.font.name = FUENTE
            t.font.size = Pt(11)
            t.font.bold = nivel == 1
            f = t.paragraph_format
            f.alignment = WD_ALIGN_PARAGRAPH.LEFT
            f.line_spacing_rule = WD_LINE_SPACING.SINGLE
            f.left_indent = Cm(0.6 * (nivel - 1))
            f.space_before = Pt(10 if nivel == 1 else 2)
            f.space_after = Pt(2)

        # Viñetas: la de la plantilla, en Calibri y sin la sangría de primera línea.
        lista = st["List Bullet"]
        lista.font.name = FUENTE
        lista.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        numerada = st["List Number"]
        numerada.font.name = FUENTE
        numerada.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        # Dentro de las tablas, interlineado simple (APA lo permite) y a la izquierda.
        celda = st.add_style("Celda", WD_STYLE_TYPE.PARAGRAPH)
        celda.base_style = normal
        celda.font.size = Pt(10)
        celda.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        celda.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        celda.paragraph_format.space_before = Pt(2)
        celda.paragraph_format.space_after = Pt(2)

        # Rótulo de figuras y tablas: "Figura 3" en negrita, el título en cursiva abajo.
        rotulo = st.add_style("Rotulo", WD_STYLE_TYPE.PARAGRAPH)
        rotulo.base_style = normal
        rotulo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rotulo.paragraph_format.keep_with_next = True
        rotulo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

        nota = st.add_style("Nota", WD_STYLE_TYPE.PARAGRAPH)
        nota.base_style = normal
        nota.font.size = Pt(10)
        nota.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        nota.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        nota.paragraph_format.space_before = Pt(4)
        nota.paragraph_format.space_after = Pt(12)

    def _geometria(self, s):
        s.page_width = ANCHO_PAGINA
        s.page_height = ALTO_PAGINA
        s.left_margin = MARGEN_INTERIOR   # con mirrorMargins, "left" es el interior
        s.right_margin = MARGEN_EXTERIOR
        s.top_margin = MARGEN_VERTICAL
        s.bottom_margin = MARGEN_VERTICAL
        s.header_distance = Cm(1.25)
        s.footer_distance = Cm(1.25)
        s.gutter = Cm(0)

    # -- encabezados y pies --------------------------------------------------

    def _numero_de_pagina(self, pie, alineacion):
        p = pie.paragraphs[0]
        p.text = ""
        p.style = self.doc.styles["Normal"]
        p.paragraph_format.alignment = alineacion
        p.paragraph_format.first_line_indent = Cm(0)
        _campo(p, "PAGE", "1")
        for r in p.runs:
            r.font.size = Pt(10)

    def _encabezado(self, cabecera, texto, alineacion):
        p = cabecera.paragraphs[0]
        p.text = ""
        p.paragraph_format.alignment = alineacion
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        r = p.add_run(texto)
        r.font.size = Pt(9)
        r.font.color.rgb = GRIS
        _borde_inferior(p, "A6A6A6")

    def _vaciar(self, parte):
        parte.is_linked_to_previous = False
        for p in parte.paragraphs:
            p.text = ""

    def _cabeceras(self, s, encabezado_impar, numeracion="decimal"):
        """Doble faz: el número va en la esquina exterior; el encabezado de las
        impares dice la parte, el de las pares el título del libro. La primera
        página de cada sección (la portadilla) va limpia."""
        s.different_first_page_header_footer = True
        self._vaciar(s.first_page_header)
        self._vaciar(s.first_page_footer)

        s.header.is_linked_to_previous = False
        s.even_page_header.is_linked_to_previous = False
        s.footer.is_linked_to_previous = False
        s.even_page_footer.is_linked_to_previous = False
        if encabezado_impar:
            self._encabezado(s.header, encabezado_impar, WD_ALIGN_PARAGRAPH.RIGHT)
            self._encabezado(s.even_page_header, TITULO_CORTO, WD_ALIGN_PARAGRAPH.LEFT)
        else:
            self._vaciar(s.header)
            self._vaciar(s.even_page_header)
        self._numero_de_pagina(s.footer, WD_ALIGN_PARAGRAPH.RIGHT)
        self._numero_de_pagina(s.even_page_footer, WD_ALIGN_PARAGRAPH.LEFT)

    def nueva_seccion_impar(self):
        s = self.doc.add_section(WD_SECTION.ODD_PAGE)
        self._geometria(s)
        # La sección nueva copia la anterior, incluido el "empezar en 1": se le saca
        # para que la numeración siga corrida.
        _numeracion(s, "decimal")
        return s

    # -- bloques de contenido -------------------------------------------------

    def salto(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def parte(self, numero, nombre):
        """Portadilla de la parte, sola en una página impar."""
        s = self.nueva_seccion_impar()
        self._cabeceras(s, f"{numero}. {nombre}")
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Cm(8)
        r = p.add_run(f"Parte {numero}")
        r.font.size = Pt(14)
        r.font.color.rgb = GRIS
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h = self.doc.add_heading(f"{numero}. {nombre}", level=1)
        _borde_inferior(h, "000000")
        self.salto()

    def titulo(self, texto, nivel=1):
        """Niveles APA: 1 centrado, 2 izquierda, 3 izquierda en cursiva."""
        return self.doc.add_heading(texto, level=nivel + 1)

    def parrafo(self, texto, sangria=True, junto=False):
        """junto=True: no se separa de lo que sigue (por ejemplo, de su tabla)."""
        p = self.doc.add_paragraph(style="Cuerpo" if sangria else "Normal")
        p.paragraph_format.keep_with_next = junto or None
        return escribir_rico(p, texto)

    def vinetas(self, items):
        for it in items:
            escribir_rico(self.doc.add_paragraph(style="List Bullet"), it)

    def numerada(self, items):
        # Cada lista reinicia su numeración: se le da un numId propio.
        numbering = self.doc.part.numbering_part.numbering_definitions._numbering
        estilo = self.doc.styles["List Number"]
        abstract = None
        num_id_base = estilo.element.pPr.numPr.numId.val
        for n in numbering.findall(qn("w:num")):
            if n.get(qn("w:numId")) == str(num_id_base):
                abstract = n.find(qn("w:abstractNumId")).get(qn("w:val"))
        nuevo = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
        num = _el("w:num", **{"w:numId": nuevo})
        num.append(_el("w:abstractNumId", **{"w:val": abstract}))
        ov = _el("w:lvlOverride", **{"w:ilvl": 0})
        ov.append(_el("w:startOverride", **{"w:val": 1}))
        num.append(ov)
        numbering.append(num)
        for it in items:
            p = self.doc.add_paragraph(style="List Number")
            ppr = p._p.get_or_add_pPr()
            numpr = _el("w:numPr")
            numpr.append(_el("w:ilvl", **{"w:val": 0}))
            numpr.append(_el("w:numId", **{"w:val": nuevo}))
            ppr.append(numpr)
            escribir_rico(p, it)

    def _rotulo(self, tipo, numero, titulo):
        p = self.doc.add_paragraph(style="Rotulo")
        p.paragraph_format.space_before = Pt(12)
        p.add_run(f"{tipo} {numero}").bold = True
        p = self.doc.add_paragraph(style="Rotulo")
        p.add_run(titulo).italic = True

    def _nota(self, nota):
        if nota:
            p = self.doc.add_paragraph(style="Nota")
            p.add_run("Nota. ").italic = True
            escribir_rico(p, nota)
        else:
            self.doc.add_paragraph(style="Nota")

    def figura(self, ruta, titulo, nota=None, ancho_cm=15.0):
        self.figuras += 1
        self._rotulo("Figura", self.figuras, titulo)
        p = self.doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.add_run().add_picture(str(ruta), width=Cm(min(ancho_cm, 15.0)))
        self._nota(nota)

    def tabla(self, encabezados, filas, anchos_cm, titulo=None, nota=None, tamano=10):
        """Tabla APA: solo líneas horizontales, encabezado que se repite en cada
        página y filas que no se parten entre dos páginas."""
        if titulo:
            self.tablas += 1
            self._rotulo("Tabla", self.tablas, titulo)
        t = self.doc.add_table(rows=1, cols=len(encabezados))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblpr = t._tbl.tblPr
        bordes = _el("w:tblBorders")
        for lado, sz in (("top", 8), ("bottom", 8), ("insideH", 2)):
            bordes.append(_el(f"w:{lado}", **{"w:val": "single", "w:sz": sz, "w:space": 0, "w:color": "000000" if lado != "insideH" else "BFBFBF"}))
        for lado in ("left", "right", "insideV"):
            bordes.append(_el(f"w:{lado}", **{"w:val": "nil"}))
        tblpr.append(bordes)
        tblpr.append(_el("w:tblW", **{"w:w": int(sum(anchos_cm) * 567), "w:type": "dxa"}))
        grid = t._tbl.tblGrid
        for gc, a in zip(grid.findall(qn("w:gridCol")), anchos_cm):
            gc.set(qn("w:w"), str(int(a * 567)))

        def llenar(fila, valores, encabezado=False):
            trpr = fila._tr.get_or_add_trPr()
            trpr.append(_el("w:cantSplit"))
            if encabezado:
                trpr.append(_el("w:tblHeader"))
            for celda, valor, ancho in zip(fila.cells, valores, anchos_cm):
                celda.width = Cm(ancho)
                p = celda.paragraphs[0]
                p.style = self.doc.styles["Celda"]
                # El encabezado nunca queda solo al pie de una página.
                if encabezado:
                    p.paragraph_format.keep_with_next = True
                escribir_rico(p, "" if valor is None else str(valor), tamano)
                if encabezado:
                    for r in p.runs:
                        r.bold = True
                    tcpr = celda._tc.get_or_add_tcPr()
                    tcpr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": "E7E6E6"}))
                    tcb = _el("w:tcBorders")
                    tcb.append(_el("w:bottom", **{"w:val": "single", "w:sz": 6, "w:space": 0, "w:color": "000000"}))
                    tcpr.append(tcb)

        llenar(t.rows[0], encabezados, encabezado=True)
        for valores in filas:
            llenar(t.add_row(), valores)
        self._nota(nota)
        return t

    def pendiente(self, texto):
        """Marca de contenido que llega en una fase posterior. Gris y en cursiva,
        para que no se confunda con texto del libro."""
        p = self.doc.add_paragraph(style="Normal")
        r = p.add_run(f"[{texto}]")
        r.italic = True
        r.font.color.rgb = GRIS

    # -- piezas fijas --------------------------------------------------------

    def portada(self):
        """Portada de estudiante APA 7. Sin encabezado ni número."""
        s = self.doc.sections[0]
        s.different_first_page_header_footer = True
        for parte in (s.header, s.footer, s.even_page_header, s.even_page_footer,
                      s.first_page_header, s.first_page_footer):
            self._vaciar(parte)

        def linea(texto, tam=11, negrita=False, cursiva=False, antes=0, color=None):
            p = self.doc.add_paragraph(style="Normal")
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(antes)
            r = p.add_run(texto)
            r.font.size = Pt(tam)
            r.bold = negrita
            r.italic = cursiva
            if color:
                r.font.color.rgb = color
            return p

        p = self.doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Cm(2)
        p.add_run().add_picture(str(RAIZ / "docs/branding/logos/logo-icono.png"), width=Cm(5))

        linea("Sistema de Gestión Integral para La Juanita Studio", 18, negrita=True, antes=36)
        linea("Sello discográfico y estudio de formación musical", 12, cursiva=True, antes=6, color=GRIS)

        linea("Ignacio Lawson", 11, antes=96)
        linea("Ingeniería Informática, Universidad del Salvador — Sede Pilar")
        linea("Seminario de Integración Profesional")
        linea("Prof. Azul de León Aboy, Prof. Martín Machain y Prof. Ciro Nievas")
        linea("Diciembre de 2026")

    def indice(self):
        s = self.nueva_seccion_impar()
        self._cabeceras(s, None)
        s.different_first_page_header_footer = False
        _numeracion(s, "lowerRoman", 1)
        # Con la cara de un título APA de nivel 1, pero sin ser un título: así no
        # entra al índice.
        h = self.doc.add_paragraph(style="Normal")
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.space_after = Pt(12)
        h.add_run("Índice").bold = True
        p = self.doc.add_paragraph(style="Normal")
        _campo(p, 'TOC \\o "1-3" \\h \\z \\u', "Abrí el documento en Word o corré exportar.ps1 para generar el índice.")


def armar():
    libro = Libro()
    libro.portada()
    libro.indice()
    for i, (numero, nombre, modulo) in enumerate(PARTES):
        libro.parte(numero, nombre)
        if i == 0:
            _numeracion(libro.doc.sections[-1], "decimal", 1)
        importlib.import_module(f"partes.{modulo}").escribir(libro)
    # Referencias: sin portadilla. Va como Heading 1 para quedar al primer nivel del
    # índice, pero con la cara de un título APA de nivel 1 (centrado, 11 pt).
    s = libro.nueva_seccion_impar()
    libro._cabeceras(s, "Referencias")
    s.different_first_page_header_footer = False
    h = libro.doc.add_heading("Referencias", level=1)
    for r in h.runs:
        r.font.size = Pt(11)
    h.paragraph_format.space_after = Pt(12)
    importlib.import_module("partes.referencias").escribir(libro)

    libro.doc.core_properties.author = "Ignacio Lawson"
    libro.doc.core_properties.title = "Sistema de Gestión Integral para La Juanita Studio"
    libro.doc.save(SALIDA)
    print(f"{SALIDA}  ·  {libro.figuras} figuras · {libro.tablas} tablas")


if __name__ == "__main__":
    import sys
    sys.path.insert(0, str(AQUI))
    armar()
