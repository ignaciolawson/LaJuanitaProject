"""Generador del Manual de Usuario: un documento aparte del libro, con el mismo formato.

    python docs/EntregaFinal/fuente/manual.py
    pwsh docs/EntregaFinal/fuente/exportar.ps1 Manual-de-Usuario

Las capturas se declaran con `cap()` dentro del contenido. Si el archivo existe en
fuente/capturas/<id>.png entra al documento; si no, queda un recuadro que dice cuál
falta. La misma lista genera docs/EntregaFinal/capturas.md, la guía para sacarlas.
"""

import sys
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))

from libro import GRIS, Libro, _el, escribir_rico  # noqa: E402

SALIDA = AQUI.parent / "Manual-de-Usuario.docx"
GUIA = AQUI.parent / "capturas.md"
CAPTURAS = AQUI / "capturas"


class Manual(Libro):
    def __init__(self):
        super().__init__("La Juanita Studio · Manual de Usuario")
        self.capturas = []
        self.capitulo_actual = ""

    def capitulo(self, numero, nombre):
        """Un capítulo arranca en página impar, con el título arriba y sin portadilla:
        en un manual, una página casi vacía por capítulo es papel que no ayuda."""
        s = self.nueva_seccion_impar()
        self._cabeceras(s, f"{numero}. {nombre}")
        s.different_first_page_header_footer = False
        self.capitulo_actual = f"{numero}. {nombre}"
        h = self.doc.add_heading(f"{numero}. {nombre}", level=1)
        for r in h.runs:
            r.font.size = Pt(20)
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h.paragraph_format.space_after = Pt(18)

    def recuadro(self, etiqueta, texto):
        """Un aviso destacado: «Importante», «Consejo». Sombreado gris claro con una
        línea a la izquierda, que se lee igual impreso en blanco y negro."""
        p = self.doc.add_paragraph(style="Normal")
        ppr = p._p.get_or_add_pPr()
        bdr = _el("w:pBdr")
        bdr.append(_el("w:left", **{"w:val": "single", "w:sz": 24, "w:space": 8,
                                    "w:color": "404040"}))
        ppr.append(bdr)
        ppr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": "F2F2F2"}))
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(10)
        p.add_run(f"{etiqueta}. ").bold = True
        escribir_rico(p, texto)

    def cap(self, id_, titulo, ruta, usuario, preparar, mostrar, ancho_cm=15.0):
        """Una captura de pantalla. Todo lo que hace falta para sacarla queda escrito acá
        y pasa tal cual a capturas.md."""
        self.capturas.append({
            "id": id_, "titulo": titulo, "ruta": ruta, "usuario": usuario,
            "preparar": preparar, "mostrar": mostrar, "capitulo": self.capitulo_actual,
            "figura": self.figuras + 1,
        })
        archivo = CAPTURAS / f"{id_}.png"
        if archivo.exists():
            self.figura(archivo, titulo, ancho_cm=ancho_cm)
            return
        self.figuras += 1
        self._rotulo("Figura", self.figuras, titulo)
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        celda = t.rows[0].cells[0]
        celda.width = Cm(min(ancho_cm, 15.0))
        tr = t.rows[0]._tr.get_or_add_trPr()
        tr.append(_el("w:trHeight", **{"w:val": int(5.5 * 567), "w:hRule": "atLeast"}))
        tcpr = celda._tc.get_or_add_tcPr()
        tcb = _el("w:tcBorders")
        for lado in ("top", "left", "bottom", "right"):
            tcb.append(_el(f"w:{lado}", **{"w:val": "dashed", "w:sz": 8, "w:space": 0,
                                           "w:color": "808080"}))
        tcpr.append(tcb)
        tcpr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": "F7F7F7"}))
        tcpr.append(_el("w:vAlign", **{"w:val": "center"}))
        p = celda.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"CAPTURA PENDIENTE · {id_}")
        r.bold = True
        r.font.color.rgb = GRIS
        p2 = celda.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p2.add_run(f"Guardar como fuente/capturas/{id_}.png — ver capturas.md")
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        self.doc.add_paragraph(style="Nota")


def _guia(manual):
    hechas = sum((CAPTURAS / f"{c['id']}.png").exists() for c in manual.capturas)
    lineas = [
        "# Capturas del Manual de Usuario",
        "",
        f"Son **{len(manual.capturas)} capturas**; hay {hechas} cargadas. Esta lista la genera "
        "`fuente/manual.py`: no la edites a mano, se pisa.",
        "",
        "## Cómo sacarlas",
        "",
        "1. **Datos de prueba, no reales.** El manual se imprime y se entrega: que no aparezcan "
        "nombres, teléfonos ni emails de alumnos de verdad. Revisá que la base de desarrollo no tenga "
        "datos reales cargados, y si los tiene, creá personas inventadas para las capturas.",
        "2. **Dónde abrirlo:** con el sistema levantado en tu máquina, cada ruta de la lista va "
        "después de `http://localhost:5173` (por ejemplo, `http://localhost:5173/app/login`).",
        "3. **Tema claro**, zoom del navegador al **100 %** y ventana de unos **1440 × 900**.",
        "4. **Solo la ventana del sistema**: sin la barra del navegador ni la barra de tareas "
        "(en Windows, `Win + Shift + S` y recortás).",
        "5. Las que dicen **celular**: en Chrome, `F12` → ícono de dispositivo → *iPhone 12 Pro* "
        "(390 px de ancho).",
        "6. Guardala en `docs/EntregaFinal/fuente/capturas/` con el **nombre exacto** del "
        "título (por ejemplo `02-login.png`) y regenerá el manual: entra sola en su lugar.",
        "",
        "## La lista",
        "",
    ]
    capitulo = None
    for c in manual.capturas:
        if c["capitulo"] != capitulo:
            capitulo = c["capitulo"]
            lineas += [f"### {capitulo}", ""]
        hecha = "x" if (CAPTURAS / f"{c['id']}.png").exists() else " "
        lineas += [
            f"- [{hecha}] **`{c['id']}.png`** — Figura {c['figura']}: *{c['titulo']}*",
            f"  - **Dónde:** `{c['ruta']}` · **Entrar como:** {c['usuario']}",
            f"  - **Antes:** {c['preparar']}",
            f"  - **Que se vea:** {c['mostrar']}",
            "",
        ]
    GUIA.write_text("\n".join(lineas), encoding="utf-8")


def armar():
    import manual_contenido
    m = Manual()
    m.portada("Manual de Usuario", "Sistema de Gestión Integral para La Juanita Studio")
    m.indice()
    manual_contenido.escribir(m)
    m.doc.core_properties.author = "Ignacio Lawson"
    m.doc.core_properties.title = "Manual de Usuario — La Juanita Studio"
    m.doc.save(SALIDA)
    _guia(m)
    print(f"{SALIDA}  ·  {len(m.capturas)} capturas · {m.tablas} tablas")


if __name__ == "__main__":
    armar()
