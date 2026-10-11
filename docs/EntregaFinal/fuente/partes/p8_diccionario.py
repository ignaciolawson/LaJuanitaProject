"""El diccionario de datos completo, leído de docs/db/diccionario-de-datos.xlsx.
Si el diccionario cambia, se vuelve a correr libro.py y entra solo."""

from collections import OrderedDict

import openpyxl

from libro import RAIZ


def _leer():
    wb = openpyxl.load_workbook(RAIZ / "docs/db/diccionario-de-datos.xlsx", read_only=True)
    descripciones = {r[0]: r[2] for r in wb["Tablas"].iter_rows(min_row=2, values_only=True) if r[0]}
    tablas = OrderedDict()
    for r in wb["Diccionario"].iter_rows(min_row=4, values_only=True):
        if r[0]:
            tablas.setdefault(r[0], []).append(r[1:6])
    return descripciones, tablas


def escribir(libro):
    descripciones, tablas = _leer()
    total = sum(len(c) for c in tablas.values())
    libro.titulo("Diccionario de datos")
    libro.parrafo(
        f"A continuación se describen las {len(tablas)} tablas de la base de datos y sus "
        f"{total} columnas, en orden alfabético. Para cada columna se indica su tipo de dato, "
        "su tamaño, si es obligatoria y qué guarda. Las relaciones entre tablas se ven en las "
        "láminas de la Parte V.")
    for nombre, columnas in tablas.items():
        libro.titulo(nombre, nivel=3)
        if descripciones.get(nombre):
            libro.parrafo(descripciones[nombre], sangria=False, junto=True)
        libro.tabla(["Campo", "Tipo de dato", "Tamaño", "Req.", "Descripción"],
                    columnas, [3.7, 2.8, 1.9, 1.0, 5.6], tamano=9)
