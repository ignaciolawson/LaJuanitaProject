"""Parte VII — Diagrama de Gantt, ajustado a las fechas reales (2026-10-10).

S1 es la semana del 1 de abril de 2026; la presentación final (16 de diciembre) cae en
la S38. Hasta agosto no hay repositorio, así que esas semanas salen de las fechas de los
documentos: la entrevista (17/4, S3), el relevamiento V2 (5/5, S5), la propuesta
comercial (16/6, S11) y la pre-entrega (5/8, S19). Desde el 6/8 salen del historial de
git. Las horas son las estimadas, como en la planilla."""

import datetime as dt

from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from libro import ALTO_PAGINA, ANCHO_PAGINA, _el, escribir_rico

INICIO = dt.date(2026, 4, 1)
SEMANAS = 38

# (grupo, [(actividad, horas estimadas, [(desde, hasta), ...])])
GRUPOS = [
    ("Inicio y análisis", [
        ("Estudio de prefactibilidad", 2, [(1, 1)]),
        ("Relevamiento", 8, [(1, 3)]),
        ("Redacción del informe de relevamiento", 13, [(3, 4)]),
        ("Revisión del informe de relevamiento", 2, [(5, 5)]),
        ("Hoja de horas", 1, [(5, 5)]),
        ("Diagrama de Gantt", 1, [(5, 5)]),
    ]),
    ("Propuesta", [
        ("Propuesta técnica", 20, [(6, 9)]),
        ("Propuesta comercial", 10, [(9, 11)]),
        ("Revisión de las propuestas", 3, [(11, 12)]),
    ]),
    ("Diseño y base de datos", [
        ("Elaboración de diagramas", 15, [(5, 5), (28, 28)]),
        ("Implementación de la base de datos", 10, [(19, 19)]),
    ]),
    ("Desarrollo", [
        ("Módulo 1 – Alumnos", 48, [(20, 20)]),
        ("Módulo 2 – Horarios y salas", 24, [(20, 20)]),
        ("Módulo 3 – Pagos", 15, [(20, 20)]),
        ("Módulo 4 – Portal del alumno", 72, [(21, 21)]),
        ("Módulo 5 – Portal del profesor", 72, [(21, 21)]),
        ("Módulo 6 – Mix & Mastering", 48, [(21, 21)]),
        ("Módulo 7 – Sello", 24, [(21, 21)]),
        ("Módulo 8 – Tablero", 48, [(21, 21)]),
    ]),
    ("Pruebas y documentación", [
        ("Pruebas funcionales", 48, [(21, 26)]),
        ("Manual de usuario", 15, [(30, 34)]),
    ]),
    ("Cierre", [
        ("Presentación final", 2, [(38, 38)]),
    ]),
]

# Tonos de gris por grupo: se distinguen impresos en blanco y negro.
TONO = ["BFBFBF", "A6A6A6", "7F7F7F", "595959", "8C8C8C", "262626"]
BANDA = "E7E6E6"
LINEA = "BFBFBF"

ANCHO_ACTIVIDAD = 4.9
ANCHO_HORAS = 0.9
ANCHO_SEMANA = 0.47


def _sombra(celda, color):
    celda._tc.get_or_add_tcPr().append(
        _el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": color}))


def _texto(celda, texto, tam=7, negrita=False, centro=False):
    p = celda.paragraphs[0]
    p.style = "Celda"
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    if centro:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto)
    r.font.size = Pt(tam)
    r.bold = negrita


def _sin_saltos(fila):
    trpr = fila._tr.get_or_add_trPr()
    trpr.append(_el("w:cantSplit"))


def _apaisada(libro):
    s = libro.doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = ALTO_PAGINA, ANCHO_PAGINA
    s.left_margin = s.right_margin = Cm(2.5)
    s.top_margin = Cm(3.5)       # el lado largo de arriba es el que va al lomo
    s.bottom_margin = Cm(2.0)
    s.different_first_page_header_footer = False
    return s


def _volver_a_vertical(libro):
    s = libro.doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.PORTRAIT
    libro._geometria(s)
    return s


def escribir(libro):
    libro.titulo("Diagrama de Gantt")
    libro.parrafo(
        "El diagrama muestra en qué semanas se realizó cada actividad del proyecto, desde el "
        "1 de abril de 2026 (semana 1) hasta la presentación final, el 16 de diciembre "
        "(semana 38). La columna de horas indica el tiempo estimado de cada actividad; el "
        "tiempo real está en la Parte VI. Las actividades de la primera etapa se ubican por "
        "la fecha de sus documentos y las del desarrollo, por el historial del repositorio.")

    _apaisada(libro)
    libro._rotulo("Figura", libro.figuras + 1, "Diagrama de Gantt del proyecto")
    libro.figuras += 1

    columnas = 2 + SEMANAS
    t = libro.doc.add_table(rows=0, cols=columnas)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    anchos = [ANCHO_ACTIVIDAD, ANCHO_HORAS] + [ANCHO_SEMANA] * SEMANAS
    tblpr = t._tbl.tblPr
    bordes = _el("w:tblBorders")
    for lado in ("top", "bottom", "left", "right", "insideH", "insideV"):
        bordes.append(_el(f"w:{lado}", **{"w:val": "single", "w:sz": 2, "w:space": 0,
                                          "w:color": LINEA}))
    tblpr.append(bordes)
    tblpr.append(_el("w:tblW", **{"w:w": int(sum(anchos) * 567), "w:type": "dxa"}))
    # Sin márgenes laterales en las celdas: a 0,47 cm por semana no sobra nada.
    mar = _el("w:tblCellMar")
    for lado in ("left", "right"):
        mar.append(_el(f"w:{lado}", **{"w:w": 20, "w:type": "dxa"}))
    tblpr.append(mar)
    for gc, a in zip(t._tbl.tblGrid.findall("{http://schemas.openxmlformats.org/"
                                            "wordprocessingml/2006/main}gridCol"), anchos):
        gc.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w",
               str(int(a * 567)))

    def fila_nueva(encabezado=False):
        f = t.add_row()
        _sin_saltos(f)
        if encabezado:
            f._tr.get_or_add_trPr().append(_el("w:tblHeader"))
        for c, a in zip(f.cells, anchos):
            c.width = Cm(a)
            # Las celdas vacías también, o heredan el 1,5 del texto y la fila crece.
            par = c.paragraphs[0]
            par.style = "Celda"
            par.paragraph_format.space_before = Pt(1)
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.line_spacing = 1.0
            fmt = par._p.get_or_add_pPr()
            rpr = _el("w:rPr")
            rpr.append(_el("w:sz", **{"w:val": 14}))
            fmt.append(rpr)
        return f

    # Fila 1: los meses, agrupando sus semanas.
    f = fila_nueva(True)
    _texto(f.cells[0], "Actividad", 8, True)
    _texto(f.cells[1], "Hs", 8, True, True)
    meses = {}
    for s in range(1, SEMANAS + 1):
        inicio = INICIO + dt.timedelta(days=7 * (s - 1))
        meses.setdefault((inicio.year, inicio.month), []).append(s)
    nombres = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov",
               "Dic"]
    for (_, mes), semanas in meses.items():
        a, b = semanas[0] + 1, semanas[-1] + 1
        celda = f.cells[a].merge(f.cells[b]) if b > a else f.cells[a]
        _texto(celda, nombres[mes - 1], 7, True, True)
    for c in f.cells:
        _sombra(c, BANDA)

    # Fila 2: los números de semana.
    f = fila_nueva(True)
    for s in range(1, SEMANAS + 1):
        _texto(f.cells[s + 1], str(s), 6, False, True)
    for c in f.cells:
        _sombra(c, BANDA)
    # Actividad y Hs ocupan las dos filas del encabezado.
    t.rows[0].cells[0].merge(f.cells[0])
    t.rows[0].cells[1].merge(f.cells[1])

    for g, (grupo, actividades) in enumerate(GRUPOS):
        f = fila_nueva()
        celda = f.cells[0].merge(f.cells[-1])
        _texto(celda, grupo.upper(), 7, True)
        _sombra(celda, BANDA)
        for nombre, horas, tramos in actividades:
            f = fila_nueva()
            _texto(f.cells[0], nombre, 7)
            _texto(f.cells[1], str(horas), 7, False, True)
            for desde, hasta in tramos:
                for s in range(desde, hasta + 1):
                    _sombra(f.cells[s + 1], TONO[g])

    libro._nota("Cada semana se cuenta desde un miércoles: la semana 1 empieza el 1 de abril "
                "de 2026. La "
                "elaboración de diagramas tiene dos tramos: los del relevamiento (semana 5) y "
                "los diagramas formales del sistema (semana 28). El manual de usuario está "
                "planificado.")
